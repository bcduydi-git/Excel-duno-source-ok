"""Kiểm thử công thức báo cáo bằng LibreOffice: đổi tham số, tính lại, so từng ô với expected.py."""
import datetime as dt
import os
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

import expected

SHEET = {"tq": ("1_TONG_QUAN", 12, 12), "hd": ("2_HDTD_HAN_MUC", 9, 13), "ku": ("3_KHE_UOC", 9, 15), "lich": ("4_LICH_TRA_NO", 10, 15)}


def chay(xlsx, scen, tables, R, hom_nay):
    wb = openpyxl.load_workbook(xlsx)
    for key, (sh, r0, nc) in SHEET.items():
        ws = wb[sh]
        for cell, v in scen.get(key, {}).items():
            ws[cell] = v
        c = ws.cell(r0, 1)
        if isinstance(c.value, ArrayFormula):
            c.value = ArrayFormula(f"A{r0}:{openpyxl.utils.get_column_letter(nc)}{r0 + 400}", c.value.text)
    d = tempfile.mkdtemp()
    p = os.path.join(d, "t.xlsx")
    wb.save(p)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", d + "/o", p],
                   capture_output=True, timeout=300)
    return openpyxl.load_workbook(d + "/o/t.xlsx", data_only=True)


def so(a, b):
    if isinstance(b, dt.date):
        b = dt.datetime(b.year, b.month, b.day)
    if b in ("", None):
        return a in ("", None)
    if isinstance(b, (int, float)) and not isinstance(b, bool):
        return isinstance(a, (int, float)) and abs(a - b) <= 0.5 + 1e-9 * abs(b)
    return a == b


def kiem(wb, key, exp, nhan):
    sh, r0, nc = SHEET[key]
    ws = wb[sh]
    loi = 0
    for i, row in enumerate(exp):
        for j, v in enumerate(row):
            a = ws.cell(r0 + i, j + 1).value
            if not so(a, v):
                loi += 1
                if loi <= 5:
                    print(f"  ✗ {nhan} {sh} R{r0 + i}C{j + 1}: LO={a!r} PY={v!r}")
    extra = ws.cell(r0 + len(exp), 1).value
    if len(exp) == 1 and len(exp[0]) == 1:
        extra = None  # chỉ có câu thông báo: LibreOffice lặp giá trị đơn ra cả vùng mảng cố định
    if extra not in (None, "", "#N/A") and not str(extra).startswith("#"):
        loi += 1
        print(f"  ✗ {nhan} {sh}: thừa dòng {extra!r}")
    print(f"  {'OK' if not loi else 'LỖI'} {nhan} {sh}: {len(exp)} dòng, {loi} ô lệch")
    return loi
