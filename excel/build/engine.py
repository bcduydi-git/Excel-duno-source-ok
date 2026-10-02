"""Bản tham chiếu bằng Python của các query Power Query trong power_query/*.m.

Dùng để: (1) tạo dữ liệu cache cho các bảng KQ_* khi dựng workbook,
(2) kiểm chứng công thức báo cáo (so với LibreOffice) và đối chiếu với logic Power BI cũ.
Mỗi hàm ở đây tương ứng một query cùng tên; giữ hai bên cùng logic khi sửa.
"""
import datetime as dt
import re
import unicodedata
from collections import defaultdict
from dateutil.relativedelta import relativedelta

import openpyxl

from source_data import as_date, as_num

D1900 = dt.date(1900, 1, 1)


# ---------------------------------------------------------------- tham số
def tham_so(tables):
    p = {r["ThamSo"]: r["GiaTri"] for r in tables["tblThamSo"]["rows"] if r.get("ThamSo")}
    return {
        "MocTheoDoiLichDaiHan": as_date(p["MocTheoDoiLichDaiHan"]),
        "NguongLechLai": as_num(p["NguongLechLai"]),
    }


# ---------------------------------------------------------------- sao kê
def _chu(x):
    return "" if x is None else str(x).strip()


def _chuan(x):
    return _chu(x).replace("\r", " ").replace("\n", " ").strip().upper()


def _so_vpb(x):
    if isinstance(x, (int, float)):
        return float(x)
    s = _chu(x).replace(" ", "").replace("\xa0", "")
    if s == "":
        return 0.0
    dc, dp = s.rfind("."), s.rfind(",")
    if dc >= 0 and dp >= 0:
        vi = dp > dc
    elif dc >= 0:
        vi = s.count(".") > 1 or len(s.split(".")[-1]) == 3
    elif dp >= 0:
        vi = not (s.count(",") > 1 or len(s.split(",")[-1]) == 3)
    else:
        vi = False
    if vi:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    return float(s)


def _ngay(x):
    if x is None or _chu(x) == "":
        return None
    if isinstance(x, (dt.datetime, dt.date)):
        return as_date(x)
    return dt.datetime.strptime(_chu(x)[:10], "%d/%m/%Y").date()


def doc_vpbank(path, ten_file):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ung_vien = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(max_col=30, values_only=True)]
        vt = -1
        for i, r in enumerate(rows[:40]):
            nd = [_chuan(v) for v in r]
            if any("NỘI DUNG GIAO DỊCH" in v for v in nd) and any("SỐ TIỀN GHI NỢ" in v for v in nd):
                vt = i
                break
        if vt >= 0:
            ung_vien.append((ws.title, rows, vt))
    if not ung_vien:
        raise ValueError("Không tìm thấy tiêu đề giao dịch VPBank: " + ten_file)
    uu = [u for u in ung_vien if _chuan(u[0]) == "TỔNG HỢP"]
    _, rows, vt = (uu or ung_vien)[0]
    td = [_chuan(v) for v in rows[vt]]

    def cot(nhan):
        n = _chuan(nhan)
        for i, v in enumerate(td):
            if v.startswith(n):
                return i
        raise ValueError("Thiếu cột " + nhan)

    c = {k: cot(v) for k, v in dict(STT="STT", SoGD="Số bút toán", Ngay="Ngày giá trị", Co="Số tiền ghi có",
                                    No="Số tiền ghi nợ", ND="Nội dung giao dịch", Gio="Thời gian giao dịch").items()}
    out = []
    for i in range(vt + 1, len(rows)):
        r = rows[i]
        stt = _chu(r[c["STT"]])
        if stt == "" or not stt.isdigit():
            continue
        gio = _chu(r[c["Gio"]])
        ngay = _ngay(gio[:10]) if gio else _ngay(r[c["Ngay"]])
        out.append(dict(DongNguon=i + 1, SoGiaoDich=_chu(r[c["SoGD"]]), NgayGiaoDich=ngay,
                        GhiNo=_so_vpb(r[c["No"]]), NoiDung=_chu(r[c["ND"]]), TenFileNguon=ten_file))
    for g in out:
        ref = g["SoGiaoDich"].upper()
        bo = ref.split(";")[0]
        ma = bo[2:] if bo.startswith("PDLD") else bo
        g["MaKheUoc"] = ma if (ma.startswith("LD") and len(ma) == 12 and ma[2:].isdigit()) else None
        nd = _chuan(g["NoiDung"])
        la_goc = nd.startswith("THANH TOAN GOC") or nd.startswith("THANH TOAN NO GOC")
        la_lai = nd.startswith("THANH TOAN LAI")
        gop = "GOC VA LAI" in nd or "GỐC VÀ LÃI" in nd
        co_ma = g["MaKheUoc"] is not None
        if g["GhiNo"] <= 0:
            pl = "KHAC"
        elif gop:
            pl = "CHO_PHAN_LOAI"
        elif la_goc and co_ma:
            pl = "GOC"
        elif la_lai and co_ma:
            pl = "LAI_PHI"
        elif co_ma or la_goc or la_lai or any(k in nd for k in ("THU NO", "THU GOC", "THU LAI")):
            pl = "CHO_PHAN_LOAI"
        else:
            pl = "KHAC"
        g["PhanLoai"] = pl
    return out


NAB_GOC = ("THU VON KHE UOC", "THU VỐN KHẾ ƯỚC")
NAB_LAI = ("THU LAI KHE UOC", "THU LÃI KHẾ ƯỚC")


def _nab_loai(nd):
    if any(nd.startswith(t) for t in NAB_GOC):
        return "GOC"
    if any(nd.startswith(t) for t in NAB_LAI):
        return "LAI_PHI"
    return ""


def doc_nabank(path, ten_file):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Sao kê"]
    out = []
    for i, r in enumerate(ws.iter_rows(max_col=8, values_only=True)):
        r = list(r) + [None] * 8
        stt = _chu(r[0])
        if stt == "" or not stt.isdigit():
            continue
        out.append(dict(DongNguon=i + 1, NgayGiaoDich=_ngay(r[1]), SoGiaoDich=_chu(r[3]), NoiDung=_chu(r[4]),
                        GhiNo=as_num(r[5]), TenFileNguon=ten_file))
    for g in out:
        nd = g["NoiDung"].upper()
        loai = _nab_loai(nd)
        ma = None
        if loai:
            tt = [t for t in (NAB_GOC if loai == "GOC" else NAB_LAI) if nd.startswith(t)][0]
            sau = g["NoiDung"][len(tt):].strip()
            if sau.startswith(":"):
                sau = sau[1:].strip()
            ma = (sau + "(").split("(")[0].strip() or None
        g["MaKheUoc"] = ma
        gop = "GOC VA LAI" in nd or "GỐC VÀ LÃI" in nd
        phi_lc = "LC" in nd and ("PHI" in nd or "PHÍ" in nd)
        if g["GhiNo"] <= 0:
            pl = "KHAC"
        elif loai == "GOC" and ma:
            pl = "GOC"
        elif loai == "LAI_PHI" and ma:
            pl = "LAI_PHI"
        elif gop or phi_lc or any(k in nd for k in ("KHE UOC", "KHẾ ƯỚC", "THU NO", "THU NỢ")):
            pl = "CHO_PHAN_LOAI"
        else:
            pl = "KHAC"
        g["PhanLoai"] = pl
    return out


def sao_ke(thu_muc_files):
    """thu_muc_files: list[(MaCongTy, MaNganHang, [đường dẫn file])]."""
    out = []
    for cty, nh, files in thu_muc_files:
        for f in files:
            ten = f.split("/")[-1]
            rows = doc_vpbank(f, ten) if nh == "VPBank" else doc_nabank(f, ten)
            for g in rows:
                g.update(MaCongTy=cty, NganHangSaoKe=nh, IDGiaoDich=f"{cty}|{nh}|{ten}|{g['DongNguon']}")
                g["GocDaTra"] = g["GhiNo"] if g["PhanLoai"] == "GOC" else 0.0
                g["LaiPhiDaTra"] = g["GhiNo"] if g["PhanLoai"] == "LAI_PHI" else 0.0
            out += rows
    return out


# ---------------------------------------------------------------- danh mục khế ước
def danh_muc_khe_uoc(tables, ts):
    moc = ts["MocTheoDoiLichDaiHan"]
    ku = []
    for r in tables["tblKheUoc"]["rows"]:
        if not r.get("MaKheUoc"):
            continue
        d = {k: r.get(k) for k in ("MaCongTy", "MaNganHang", "MaKheUoc", "MaHDTD", "PhuongThucTraGoc",
                                    "TanSuatTraGoc", "TanSuatTraLai", "GhiChu")}
        for k in ("NgayNhanNo", "NgayDaoHan", "NgayTraGocDauTien", "NgayTraLaiDauTien"):
            d[k] = as_date(r.get(k))
        d["SoTienVay"] = None if r.get("SoTienVay") is None else as_num(r["SoTienVay"])
        d["LaiSuatNam"] = r.get("LaiSuatNam")
        d["CoSoTinhLai"] = r.get("CoSoTinhLai")
        d["DuNoBanDau"] = None
        d["NgayBatDauTheoDoi"] = None
        ku.append(d)
    da_co = {d["MaKheUoc"] for d in ku}
    nhom = defaultdict(list)
    for r in tables["tblLichDaiHan"]["rows"]:
        if r.get("MaKheUoc") and r["MaKheUoc"] not in da_co and as_date(r["NgayDenHan"]) >= moc:
            nhom[(r["MaCongTy"], r["MaNganHang"], r["MaKheUoc"])].append(r)
    for (cty, nh, ma), rs in nhom.items():
        rs = sorted(rs, key=lambda x: as_date(x["NgayDenHan"]))
        ku.append(dict(MaCongTy=cty, MaNganHang=nh, MaKheUoc=ma, MaHDTD=None, PhuongThucTraGoc="THEO_LICH",
                       TanSuatTraGoc=None, TanSuatTraLai=None, GhiChu=None, NgayNhanNo=None,
                       NgayDaoHan=max(as_date(x["NgayDenHan"]) for x in rs), NgayTraGocDauTien=None,
                       NgayTraLaiDauTien=None, SoTienVay=None, LaiSuatNam=None, CoSoTinhLai=None,
                       DuNoBanDau=as_num(rs[0]["DuNoTheoLichDauKy"]), NgayBatDauTheoDoi=moc))
    for d in ku:
        dai = d["NgayBatDauTheoDoi"] is not None
        d["GocBanDau"] = d["DuNoBanDau"] if dai else d["SoTienVay"]
        d["NgayBatDau"] = d["NgayBatDauTheoDoi"] if dai else d["NgayNhanNo"]
    return ku


# ---------------------------------------------------------------- khung lịch (chưa có lãi tự sinh)
BUOC = {"THANG": 1, "QUY": 3, "6_THANG": 6, "CUOI_KY": 0}


def lich_khung(tables, ku):
    ngay_cd = {r["MaNganHang"]: r.get("NgayTraLai") for r in tables["tblNganHang"]["rows"] if r.get("MaNganHang")}
    co_lich = defaultdict(list)
    for r in tables["tblLichDaiHan"]["rows"]:
        if r.get("MaKheUoc"):
            co_lich[r["MaKheUoc"]].append(r)
    lich, loi = [], []
    for k in ku:
        ma = k["MaKheUoc"]
        if ma in co_lich:
            bd = k["NgayBatDauTheoDoi"] or D1900
            rs = sorted([r for r in co_lich[ma] if as_date(r["NgayDenHan"]) >= bd],
                        key=lambda r: (as_date(r["NgayDenHan"]), r.get("KyTraNo") or 0))
            truoc = None
            for r in rs:
                goc, tong = r.get("GocPhaiTra"), r.get("TongPhaiTra")
                lich.append(dict(MaKheUoc=ma, KyTraNo=int(r["KyTraNo"]), TuNgayTinhLai=truoc,
                                 NgayDenHan=as_date(r["NgayDenHan"]), GocPhaiTra=as_num(goc),
                                 LaiPhiNhap=None if (goc is None or tong is None) else as_num(tong) - as_num(goc),
                                 NguonLich="LICH_NGAN_HANG"))
                truoc = as_date(r["NgayDenHan"])
            continue
        ly_do = None
        nhan, dao, tien = k["NgayNhanNo"], k["NgayDaoHan"], k["SoTienVay"]
        tan_suat = k["TanSuatTraLai"]
        if k["PhuongThucTraGoc"] != "CUOI_KY":
            ly_do = "Chưa có lịch ở LICH_DAI_HAN và không tự sinh được với phương thức trả gốc này"
        elif nhan is None or dao is None:
            ly_do = "Thiếu ngày nhận nợ hoặc ngày đáo hạn"
        elif dao <= nhan:
            ly_do = "Ngày đáo hạn phải sau ngày nhận nợ"
        elif tien is None or tien <= 0:
            ly_do = "Thiếu hoặc sai số tiền vay"
        elif tan_suat not in BUOC:
            ly_do = "Tần suất trả lãi không hợp lệ"
        if ly_do is None:
            buoc = BUOC[tan_suat]
            ngay_dau = k["NgayTraLaiDauTien"]
            if ngay_dau is None and buoc > 0:
                cd = ngay_cd.get(k["MaNganHang"])
                if cd:
                    d = dt.date(nhan.year, nhan.month, int(cd))
                    ngay_dau = d if d > nhan else d + relativedelta(months=1)
            if buoc > 0 and ngay_dau is None:
                ly_do = "Cần ngày trả lãi đầu tiên (KHE_UOC) hoặc ngày trả lãi cố định của ngân hàng (DANH_MUC)"
            elif buoc > 0 and ngay_dau <= nhan:
                ly_do = "Ngày trả lãi đầu tiên phải sau ngày nhận nợ"
        if ly_do:
            loi.append(dict(MaKheUoc=ma, ChiTiet=ly_do))
            continue
        cac = []
        if buoc > 0:
            i = 0
            while ngay_dau + relativedelta(months=i * buoc) < dao:
                cac.append(ngay_dau + relativedelta(months=i * buoc))
                i += 1
        cac.append(dao)
        for i, d in enumerate(cac):
            lich.append(dict(MaKheUoc=ma, KyTraNo=i + 1, TuNgayTinhLai=nhan if i == 0 else cac[i - 1], NgayDenHan=d,
                             GocPhaiTra=tien if d == dao else 0.0, LaiPhiNhap=None, NguonLich="TU_SINH"))
    info = {k["MaKheUoc"]: k for k in ku}
    lich.sort(key=lambda r: (r["MaKheUoc"], r["NgayDenHan"], r["KyTraNo"]))
    luy = defaultdict(float)
    for r in lich:
        k = info[r["MaKheUoc"]]
        luy[r["MaKheUoc"]] += r["GocPhaiTra"]
        r.update(MaCongTy=k["MaCongTy"], MaNganHang=k["MaNganHang"], GocLuyKe=luy[r["MaKheUoc"]],
                 MaKy=f'{r["MaKheUoc"]}|{r["KyTraNo"]}')
    return lich, loi


# ---------------------------------------------------------------- giao dịch
def giao_dich(tables, ku, khung, sk):
    info = {k["MaKheUoc"]: k for k in ku}
    ngoai_le = []
    tay = []
    for r in tables["tblTraNoNhapTay"]["rows"]:
        vals = [r.get(c) for c in ("NgayGiaoDich", "MaKheUoc", "GocDaTra", "LaiPhiDaTra", "GhiChu")]
        if all(v is None for v in vals):
            continue
        g = dict(Nguon="NHAP_TAY", IDGiaoDich=f'NHAP_TAY|{r["__row"]}', MaKheUoc=r.get("MaKheUoc"),
                 NgayGiaoDich=as_date(r.get("NgayGiaoDich")) if r.get("NgayGiaoDich") is not None else None,
                 GocDaTra=as_num(r.get("GocDaTra")), LaiPhiDaTra=as_num(r.get("LaiPhiDaTra")),
                 NoiDung=r.get("GhiChu"), TenFileNguon="TRA_NO_NHAP_TAY", DongNguon=r["__row"])
        if g["MaKheUoc"] is None or g["NgayGiaoDich"] is None or (r.get("GocDaTra") is None and r.get("LaiPhiDaTra") is None):
            ngoai_le.append(_nl("Thiếu thông tin ở dòng nhập tay (không tính)", g,
                                "Cần đủ ngày giao dịch, mã khế ước và số tiền"))
            continue
        if g["MaKheUoc"] not in info:
            ngoai_le.append(_nl("Mã khế ước không có trong danh mục (không tính)", g, g["NoiDung"]))
            continue
        tay.append(g)
    tu_dong = []
    for s in sk:
        g = dict(Nguon="SAO_KE", IDGiaoDich=s["IDGiaoDich"], MaKheUoc=s["MaKheUoc"], NgayGiaoDich=s["NgayGiaoDich"],
                 GocDaTra=s["GocDaTra"], LaiPhiDaTra=s["LaiPhiDaTra"], NoiDung=s["NoiDung"],
                 TenFileNguon=s["TenFileNguon"], DongNguon=s["DongNguon"], GhiNo=s["GhiNo"])
        if s["PhanLoai"] == "CHO_PHAN_LOAI":
            g["SoTienGhi"] = s["GhiNo"]
            ngoai_le.append(_nl("Sao kê liên quan vay chưa tách được gốc/lãi (nhập tay nếu cần)", g, s["NoiDung"]))
        elif s["PhanLoai"] in ("GOC", "LAI_PHI"):
            if s["MaKheUoc"] not in info:
                ngoai_le.append(_nl("Mã khế ước không có trong danh mục (không tính)", g, s["NoiDung"]))
            else:
                tu_dong.append(g)
    khoa_sk = defaultdict(list)
    for g in tu_dong:
        khoa_sk[(g["MaKheUoc"], g["NgayGiaoDich"])].append(g)
    giu = []
    for g in tay:
        trung = khoa_sk.get((g["MaKheUoc"], g["NgayGiaoDich"]))
        if trung:
            sk_goc = sum(x["GocDaTra"] for x in trung)
            sk_lai = sum(x["LaiPhiDaTra"] for x in trung)
            ngoai_le.append(_nl("Nhập tay trùng ngày với sao kê – đã bỏ dòng nhập tay", g,
                                f"Sao kê cùng ngày: gốc {sk_goc:,.0f}, lãi/phí {sk_lai:,.0f}"))
        else:
            giu.append(g)
    gd = tu_dong + giu
    lich_theo_ma = defaultdict(list)
    for r in khung:
        lich_theo_ma[r["MaKheUoc"]].append(r)
    for g in gd:
        k = info[g["MaKheUoc"]]
        g.update(MaCongTy=k["MaCongTy"], MaNganHang=k["MaNganHang"], MaHDTD=k["MaHDTD"])
        g["TinhDuNo"] = 0 if (k["NgayBatDauTheoDoi"] and g["NgayGiaoDich"] < k["NgayBatDauTheoDoi"]) else 1
        if k["NgayNhanNo"] and g["NgayGiaoDich"] < k["NgayNhanNo"]:
            ngoai_le.append(_nl("Ngày trả trước ngày nhận nợ (vẫn tính)", g, f'Nhận nợ {k["NgayNhanNo"]:%d/%m/%Y}'))
    dau = ghep_lai(gd, lich_theo_ma, ngoai_le)
    gd.sort(key=lambda g: (g["NgayGiaoDich"], g["MaKheUoc"], g["IDGiaoDich"]))
    ngay_dau = dau
    return gd, ngoai_le, ngay_dau


def ghep_lai(gd, lich_theo_ma, ngoai_le):
    """Ghép giao dịch lãi/phí vào kỳ theo NGÀY (không theo số tiền). Một "lần trả" = mọi giao dịch lãi/phí
    của cùng khế ước trong cùng ngày; mỗi lần trả thuộc đúng một kỳ:
    - lần trả đầu tiên có dữ liệu: kỳ đến hạn gần nhất trước ngày trả (kỳ đầu nếu trả trước hạn);
      các kỳ trước đó không có dữ liệu giao dịch → coi như đã trả (TruocDuLieu), như Power BI cũ;
    - lần sau: kỳ đã đến hạn chưa có lần trả nào, cũ nhất; nếu không còn kỳ như vậy thì kỳ gần ngày trả nhất
      giữa kỳ vừa ghép (trả bổ sung) và kỳ kế tiếp (trả trước hạn).
    Trả về {MaKheUoc: ngày đến hạn của kỳ có dữ liệu đầu tiên}."""
    theo_ma = defaultdict(lambda: defaultdict(list))
    for g in gd:
        g["MaKy"] = g["NgayDenHanKy"] = None
        if g["LaiPhiDaTra"] > 0 and g["TinhDuNo"] == 1:
            theo_ma[g["MaKheUoc"]][g["NgayGiaoDich"]].append(g)
    dau = {}
    for ma, theo_ngay in theo_ma.items():
        ky = sorted(lich_theo_ma.get(ma, []), key=lambda r: (r["NgayDenHan"], r["KyTraNo"]))
        if not ky:
            for ds in theo_ngay.values():
                for g in ds:
                    ngoai_le.append(_nl("Không xác định được kỳ trả lãi", g, "Khế ước chưa có lịch trả nợ"))
            continue
        i = None   # kỳ cũ nhất chưa có lần trả (None: chưa gặp lần trả nào)
        cuoi = None  # kỳ vừa ghép
        for d in sorted(theo_ngay):
            j = -1
            while j + 1 < len(ky) and ky[j + 1]["NgayDenHan"] <= d:
                j += 1
            if i is None:
                cuoi = max(j, 0)
                dau[ma] = ky[cuoi]["NgayDenHan"]
            elif i <= j:
                cuoi = i
            elif i < len(ky) and abs((d - ky[i]["NgayDenHan"]).days) < abs((d - ky[cuoi]["NgayDenHan"]).days):
                cuoi = i
            i = max(i or 0, cuoi + 1)
            for g in theo_ngay[d]:
                g["MaKy"], g["NgayDenHanKy"] = ky[cuoi]["MaKy"], ky[cuoi]["NgayDenHan"]
    return dau


def _nl(loai, g, chi_tiet):
    so = g.get("SoTienGhi")
    if so is None:
        so = (g.get("GocDaTra") or 0) + (g.get("LaiPhiDaTra") or 0)
    return dict(Loai=loai, MaKheUoc=g.get("MaKheUoc"), NgayGiaoDich=g.get("NgayGiaoDich"), SoTien=so,
                ChiTiet=chi_tiet, Nguon=g.get("Nguon"), TenFileNguon=g.get("TenFileNguon"), DongNguon=g.get("DongNguon"))


# ---------------------------------------------------------------- lịch hoàn chỉnh + khế ước
def lich_tra_no(ku, khung, gd, ngay_dau):
    info = {k["MaKheUoc"]: k for k in ku}
    tra = defaultdict(list)
    for g in gd:
        if g["TinhDuNo"] == 1 and g["GocDaTra"]:
            tra[g["MaKheUoc"]].append((g["NgayGiaoDich"], g["GocDaTra"]))
    out = []
    for r in khung:
        r = dict(r)
        k = info[r["MaKheUoc"]]
        nd = ngay_dau.get(r["MaKheUoc"])
        r["TruocDuLieu"] = 1 if (nd is not None and r["NgayDenHan"] < nd) else 0
        if r["NguonLich"] == "TU_SINH":
            lai = None
            if k["LaiSuatNam"] is not None and k["CoSoTinhLai"] in (360, 365):
                # Lãi dự kiến theo dư nợ thực tế từng ngày trong kỳ: [TuNgay, NgayDenHan)
                tu, den, goc = r["TuNgayTinhLai"], r["NgayDenHan"], k["SoTienVay"]
                moc = sorted({tu, den} | {d for d, _ in tra[r["MaKheUoc"]] if tu < d < den})
                tong = 0.0
                for a, b in zip(moc[:-1], moc[1:]):
                    du = goc - sum(s for d, s in tra[r["MaKheUoc"]] if d <= a)
                    tong += max(0.0, du) * (b - a).days
                lai = _round(tong * k["LaiSuatNam"] / k["CoSoTinhLai"])
            r["LaiPhiDuKien"] = lai
        else:
            r["LaiPhiDuKien"] = r["LaiPhiNhap"]
        r["SoNgayTinhLai"] = (r["NgayDenHan"] - r["TuNgayTinhLai"]).days if r["TuNgayTinhLai"] else None
        r["TongPhaiTra"] = None if r["LaiPhiDuKien"] is None else r["GocPhaiTra"] + r["LaiPhiDuKien"]
        out.append(r)
    out.sort(key=lambda r: (r["NgayDenHan"], r["MaCongTy"], r["MaNganHang"], r["MaKheUoc"]))
    return out


def _round(x):
    return float(int(x + 0.5)) if x >= 0 else -float(int(-x + 0.5))


def khe_uoc(tables, ku, gd):
    ten = {r["MaCongTy"]: r["TenCongTy"] for r in tables["tblCongTy"]["rows"] if r.get("MaCongTy")}
    tra = defaultdict(list)
    for g in gd:
        if g["TinhDuNo"] == 1 and g["GocDaTra"]:
            tra[g["MaKheUoc"]].append((g["NgayGiaoDich"], g["GocDaTra"]))
    out, nl = [], []
    for k in ku:
        k = dict(k)
        k["TenCongTy"] = ten.get(k["MaCongTy"])
        cum, k["NgayTatToan"] = 0.0, None
        for d, s in sorted(tra[k["MaKheUoc"]]):
            cum += s
            if k["GocBanDau"] is not None and k["NgayTatToan"] is None and cum >= k["GocBanDau"] - 0.5:
                k["NgayTatToan"] = d
        if k["GocBanDau"] is not None and cum > k["GocBanDau"] + 0.5:
            nl.append(dict(Loai="Gốc đã trả vượt số vay", MaKheUoc=k["MaKheUoc"], NgayGiaoDich=None,
                           SoTien=cum - k["GocBanDau"], ChiTiet=f'Đã trả {cum:,.0f} / vay {k["GocBanDau"]:,.0f}',
                           Nguon=None, TenFileNguon=None, DongNguon=None))
        out.append(k)
    return out, nl


def chay(tables, thu_muc_files):
    ts = tham_so(tables)
    sk = sao_ke(thu_muc_files)
    dm = danh_muc_khe_uoc(tables, ts)
    khung, loi_lich = lich_khung(tables, dm)
    gd, nl1, ngay_dau = giao_dich(tables, dm, khung, sk)
    ltn = lich_tra_no(dm, khung, gd, ngay_dau)
    ku, nl2 = khe_uoc(tables, dm, gd)
    nl0 = [dict(Loai="Không sinh được lịch trả nợ", MaKheUoc=x["MaKheUoc"], NgayGiaoDich=None, SoTien=None,
                ChiTiet=x["ChiTiet"], Nguon=None, TenFileNguon=None, DongNguon=None) for x in loi_lich]
    nl = nl0 + nl1 + nl2
    nl.sort(key=lambda x: (x["Loai"], _k(x["MaKheUoc"]), _k(x["NgayGiaoDich"]), _k(x["DongNguon"])))
    return dict(ts=ts, sk=sk, ku=ku, ltn=ltn, gd=gd, nl=nl)


def _k(v):
    """Khoá sắp xếp: null đứng trước (như Table.Sort của Power Query)."""
    return (0, "") if v is None else (1, v)
